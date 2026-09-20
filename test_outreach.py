import logging
import re
import time

from openpyxl import load_workbook

from config import TRACKER_PATH, MIN_RELEVANCE_SCORE
from linkedin.browser import create_driver
from linkedin.people import (
    search_people_for_job,
    send_linkedin_message,
    send_connection_request,
)
from messages.templates import referral_message, connection_message
from storage.excel import (
    create_tracker,
    get_contacts_for_job,
    referral_already_received,
    add_message,
    add_activity,
    mark_message_sent,
    mark_referral_requested,
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
)
logger = logging.getLogger(__name__)


def clean(v):
    return re.sub(r"\s+", " ", str(v or "")).strip()


def activity_age_days(activity):
    """Return approximate activity age. Unknown/no visible activity => None."""
    text = clean(activity).lower()

    if not text:
        return None

    if any(x in text for x in (
        "just now", "today", "active today", "recently posted",
        "recently active", "now",
    )):
        return 0

    m = re.search(r"(\d+)\s*(minute|minutes|min|mins|hour|hours|hr|hrs)", text)
    if m:
        return 0

    m = re.search(r"(\d+)\s*(day|days)", text)
    if m:
        return int(m.group(1))

    m = re.search(r"(\d+)\s*(week|weeks|wk|wks)", text)
    if m:
        return int(m.group(1)) * 7

    m = re.search(r"(\d+)\s*(month|months|mo|mos)", text)
    if m:
        return int(m.group(1)) * 30

    m = re.search(r"(\d+)\s*(year|years|yr|yrs)", text)
    if m:
        return int(m.group(1)) * 365

    # A generic "posted" without an age is not enough evidence.
    return None


def is_active(contact):
    age = activity_age_days(contact.get("activity", contact.get("Activity", "")))
    return age is not None and age <= 30


def load_one_job():
    create_tracker(TRACKER_PATH)
    wb = load_workbook(TRACKER_PATH, read_only=True)

    try:
        sheet = wb["Jobs"]
        headers = [c.value for c in sheet[1]]
        jobs = []

        for row in sheet.iter_rows(min_row=2, values_only=True):
            job = dict(zip(headers, row))
            if not clean(job.get("Job ID")):
                continue

            try:
                score = float(job.get("Relevance Score") or 0)
            except (TypeError, ValueError):
                score = 0

            if score >= MIN_RELEVANCE_SCORE:
                jobs.append(job)

        # Highest relevance first.
        jobs.sort(
            key=lambda j: float(j.get("Relevance Score") or 0),
            reverse=True,
        )
        return jobs[0] if jobs else None
    finally:
        wb.close()


def run():
    job = load_one_job()

    if not job:
        print("No previously saved job with relevance >= 60 was found.")
        return

    job_id = clean(job.get("Job ID"))
    company = clean(job.get("Company"))
    title = clean(job.get("Job Title"))
    job_url = clean(job.get("Job URL"))

    print("\nONE-JOB OUTREACH TEST")
    print("=====================")
    print(f"Job: {title}")
    print(f"Company: {company}")
    print(f"Job ID: {job_id}")
    print(f"Score: {job.get('Relevance Score')}")
    print()

    if referral_already_received(job_id):
        print("A referral is already recorded for this job. Nothing will be sent.")
        return

    driver = create_driver()

    try:
        print("Searching people for this ONE job...")
        contacts = search_people_for_job(driver, job, max_people=10)

        print(f"People discovered: {len(contacts)}")

        active = []
        for c in contacts:
            age = activity_age_days(
                c.get("activity", c.get("Activity", ""))
            )
            person = clean(c.get("person", c.get("Person", "")))
            role = clean(c.get("role", c.get("Role", "")))
            activity = clean(c.get("activity", c.get("Activity", "")))

            print(
                f"  {person} | {role} | "
                f"activity={activity or 'NO VISIBLE ACTIVITY'} | "
                f"age={age if age is not None else 'unknown'}"
            )

            if is_active(c):
                active.append((age, c))

        if not active:
            print("\nNO OUTREACH SENT.")
            print("No discovered person had a visible activity signal within 30 days.")
            return

        # Most recently active first, then existing people priority.
        active.sort(
            key=lambda item: (
                item[0],
                -(float(
                    item[1].get(
                        "referral_priority_score",
                        item[1].get("Referral Priority Score", 0),
                    ) or 0
                )),
            )
        )

        _, contact = active[0]

        person = clean(contact.get("person", contact.get("Person", "")))
        profile_url = clean(
            contact.get("linkedin_url", contact.get("LinkedIn URL", ""))
        )
        degree = clean(
            contact.get("connection_degree", contact.get("Connection Degree", ""))
        ).lower()

        if not person or not profile_url:
            print("Selected contact has no usable profile URL. Nothing sent.")
            return

        print("\nSELECTED ONE PERSON")
        print("-------------------")
        print(f"Person: {person}")
        print(f"Role: {clean(contact.get('role', contact.get('Role', '')))}")
        print(f"Activity: {clean(contact.get('activity', contact.get('Activity', '')))}")
        print(f"Connection: {degree or 'unknown'}")

        if "1st" in degree or "first" in degree:
            message_type = "Referral Request"
            message = referral_message(
                person,
                company,
                title,
                job_id,
                job_url,
                [
                    x.strip()
                    for x in clean(
                        job.get("Matched Skills", job.get("Matched Skills", ""))
                    ).split(",")
                    if x.strip()
                ],
            )

            print("\nSending ONE referral message...")
            sent = send_linkedin_message(driver, profile_url, message)

        else:
            message_type = "Connection Request"
            message = connection_message(person, company, title)

            print("\nSending ONE connection request...")
            sent = send_connection_request(driver, profile_url, message)

        add_message(
            job_id=job_id,
            company=company,
            person=person,
            message_type=message_type,
            message=message,
            resume_type=job.get("Resume Type", ""),
            resume_name=job.get("Resume Name", ""),
            resume_path=job.get("Resume Path", ""),
            job_url=job_url,
            attachment_required="No",
            status="Sent" if sent else "Failed",
        )

        if sent:
            mark_message_sent(job_id, person, message_type)
            if message_type == "Referral Request":
                mark_referral_requested(job_id, person)
            add_activity(job_id, company, person, "Message Sent", message_type)
            print("\nSUCCESS: ONE outreach action was sent.")
        else:
            add_activity(
                job_id,
                company,
                person,
                "Message Send Failed",
                message_type,
            )
            print(
                "\nFAILED: LinkedIn did not confirm the action."
                "\nNo second person will be contacted."
            )

        time.sleep(1)

    finally:
        try:
            driver.quit()
        except Exception:
            pass


if __name__ == "__main__":
    run()
