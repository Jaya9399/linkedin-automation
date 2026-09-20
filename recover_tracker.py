from pathlib import Path
import zipfile
import shutil
import tempfile
import os
from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parent
DATA = PROJECT / "data"

CORRUPTED_BACKUP = DATA / "linkedin_tracker_backup_20260909_020402.xlsx"
CURRENT_TRACKER = DATA / "linkedin_tracker.xlsx"
RECOVERED_TRACKER = DATA / "linkedin_tracker_recovered.xlsx"


def copy_backup_sheets_into_valid_workbook():
    print("Loading current valid workbook...")
    wb = load_workbook(CURRENT_TRACKER)

    print("Reading recovered worksheet XML...")
    with zipfile.ZipFile(CORRUPTED_BACKUP, "r") as source_zip:

        sheet_xml = {}

        for i in range(1, 7):
            name = f"xl/worksheets/sheet{i}.xml"
            sheet_xml[name] = source_zip.read(name)

    # Save current workbook to a temporary file first.
    temp_path = DATA / "linkedin_tracker_recovery_temp.xlsx"

    if temp_path.exists():
        temp_path.unlink()

    wb.save(temp_path)
    wb.close()

    print("Injecting recovered worksheet data...")

    # Replace the worksheet XMLs in the valid workbook.
    fd, replacement_path = tempfile.mkstemp(
        prefix="linkedin_tracker_recovery_",
        suffix=".xlsx",
        dir=DATA,
    )
    os.close(fd)

    replacement_path = Path(replacement_path)

    with zipfile.ZipFile(temp_path, "r") as original_zip:
        with zipfile.ZipFile(
            replacement_path,
            "w",
            compression=zipfile.ZIP_DEFLATED,
        ) as output_zip:

            for item in original_zip.infolist():

                if item.filename in sheet_xml:
                    output_zip.writestr(
                        item,
                        sheet_xml[item.filename],
                    )
                else:
                    output_zip.writestr(
                        item,
                        original_zip.read(item.filename),
                    )

    temp_path.unlink()

    # Validate before keeping the recovered workbook.
    print("Validating recovered workbook...")

    test_wb = load_workbook(replacement_path, read_only=True)

    expected_sheets = [
        "Dashboard",
        "Jobs",
        "Contacts",
        "Referrals",
        "Messages",
        "Activity",
    ]

    if test_wb.sheetnames != expected_sheets:
        print("ERROR: Unexpected sheets:")
        print(test_wb.sheetnames)
        test_wb.close()
        replacement_path.unlink()
        raise SystemExit(1)

    counts = {}

    for sheet_name in expected_sheets:
        ws = test_wb[sheet_name]
        counts[sheet_name] = ws.max_row
        print(f"{sheet_name}: {ws.max_row} rows")

    test_wb.close()

    # We expect the recovered operational data.
    if counts["Jobs"] < 2:
        raise SystemExit("ERROR: Jobs data was not recovered.")

    if counts["Contacts"] < 2:
        raise SystemExit("ERROR: Contacts data was not recovered.")

    if counts["Messages"] < 2:
        raise SystemExit("ERROR: Messages data was not recovered.")

    if counts["Activity"] < 2:
        raise SystemExit("ERROR: Activity data was not recovered.")

    # Move validated recovery into the requested output filename.
    if RECOVERED_TRACKER.exists():
        RECOVERED_TRACKER.unlink()

    os.replace(replacement_path, RECOVERED_TRACKER)

    print()
    print("=" * 70)
    print("RECOVERY SUCCESSFUL")
    print("=" * 70)
    print(f"Recovered workbook: {RECOVERED_TRACKER}")
    print()
    print("Recovered row counts:")
    for name, count in counts.items():
        print(f"  {name}: {count}")

    print()
    print("IMPORTANT:")
    print("The original linkedin_tracker.xlsx was NOT modified.")


if __name__ == "__main__":
    copy_backup_sheets_into_valid_workbook()
